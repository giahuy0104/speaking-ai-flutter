from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Cm, Pt, RGBColor


OUT = Path(__file__).parent
TARGET = OUT / 'HOMI_TEST_CASE_ANDROID_THEO_LUONG_2026-09-18.docx'

# Each case is one executable flow with observable acceptance criteria.
SECTIONS = [
    ('1 Thiết lập và kết nối H20',
     'Dùng hồ sơ thử nghiệm. Không gỡ app hoặc xóa dữ liệu thật chỉ để chạy các ca cài mới.', [
        ('A01', 'Cài mới và xác nhận phụ huynh',
         'Trên máy hoặc hồ sơ test mới, mở HOMI; đọc pháp lý, xác nhận phụ huynh và đồng ý dùng giọng nói.',
         'Đi đúng 3 bước thiết lập. Chưa đọc/xác nhận thì chưa mở được bước dùng giọng nói; không tự bật mic.'),
        ('A02', 'Từ chối rồi cấp lại quyền',
         'Từ chối Micro/Bluetooth, thử tiếp tục; sau đó cấp lại và kết nối H20.',
         'Báo đúng quyền thiếu; không treo hoặc báo kết nối giả. Cấp lại xong có thể tiếp tục, không phải cài lại.'),
        ('A03', 'Chọn nhóm tuổi',
         'Chọn tuổi, hoàn tất rồi mở Chủ đề. Lặp với 3–5, 6–7, 8–10, 11–12 và 13–15 tuổi.',
         'Danh mục và bài học đúng nhóm tuổi; lựa chọn được giữ khi mở lại. Trẻ không tự đổi tuổi ngoài cổng phụ huynh.'),
        ('A04', 'Kết nối đủ nút và mic',
         'Bật H20; ghép đôi và kết nối trong HOMI. Nhấn MAIN, nghe lời dẫn rồi trả lời gần mic H20.',
         'Chỉ báo hoàn tất khi cả BLE và đường âm thanh H20 sẵn sàng. Loa và mic thực tế đều là H20, không chỉ có nút MAIN.'),
        ('A05', 'Pháp lý và chế độ giới hạn',
         'Mở Chính sách quyền riêng tư, Điều khoản sử dụng, Hỗ trợ; thử “Tiếp tục không dùng giọng nói”.',
         'Đủ 3 liên kết mở được. Chế độ không giọng nói không tự xin/mở mic; có đường để phụ huynh bật lại sau.'),
        ('A06', 'Cài đè giữ dữ liệu',
         'Trên hồ sơ test đã có bài dở, ngôi sao, từ phụ huynh và lịch sử, cài đè APK cần kiểm thử rồi mở lại.',
         'Giữ nguyên tuổi, dữ liệu và tiến độ; không nhân đôi mục. Ghi đúng tên APK/phiên bản đang test vào đầu tài liệu.'),
    ]),
    ('2 MAIN và chuyển giữa các nội dung',
     'Chạy bằng nút MAIN trên màn hình và nút MAIN của H20. Chờ HOMI nói xong trước khi trả lời.', [
        ('B01', 'MAIN tại trang chủ',
         'Ở trang chủ nhấn MAIN một lần; lần lượt chọn Dịch tiếng Anh, Chủ đề và Bộ từ vựng.',
         'Đọc đủ ba hướng, mở đúng mục theo tuổi. Một lần nhấn chỉ tạo một lời dẫn và một lượt nghe.'),
        ('B02', 'Chuyển nội dung hai chiều',
         'Từ mỗi mục Dịch, Chủ đề, Từ vựng, dùng MAIN/lệnh chuyển sang từng mục còn lại. Thử đủ 6 hướng.',
         'Dừng lượt cũ trước khi vào mục mới; không dịch câu lệnh điều hướng, không còn tiếng cũ hoặc hai mic cùng chạy.'),
        ('B03', 'Chuyển vào Dịch có đủ mở đầu',
         'Đang học Chủ đề hoặc Từ vựng, yêu cầu “Dịch tiếng Anh”; sau đó nói một câu mới.',
         'Phát câu chuyển, rồi “Bạn cứ nói từng câu. Muốn dừng thì nói ‘Dừng lại’.”, sau đó mới mở mic; không bỏ intro hoặc đọc hai lần.'),
        ('B04', 'Chọn Chủ đề bằng giọng nói',
         'Vào Chủ đề qua MAIN; chọn level, chủ đề và bài bằng số/tên được giới thiệu. Thử một số không hợp lệ.',
         'Các lựa chọn đúng nhóm tuổi và đúng danh sách vừa đọc. Lệnh sai được nhắc lại có giới hạn, không tự chọn bài.'),
        ('B05', 'Chọn nhóm Từ vựng',
         'Vào Bộ từ vựng qua MAIN; lần lượt chọn Ba mẹ đã thêm, Ngôi sao của con và Luyện lại.',
         'Vào đúng nhóm và đúng dữ liệu như thao tác chạm. Nhóm rỗng có thông báo phù hợp, không mở lượt ghi vô nghĩa.'),
        ('B06', 'MAIN trùng và Back',
         'Nhấn MAIN nhanh 3 lần; sau đó thử Back điện thoại khi HOMI đang nói, đang nghe và đã im lặng trong bài học.',
         'Không lặp menu hoặc chồng tiếng. Back thoát đúng màn, dừng lượt cũ; không đơ hoặc phát lại kết quả cũ sau khi đã rời.'),
        ('B07', 'Nhận lệnh lần đầu và các lượt sau',
         'Sau khi mở app, dùng MAIN chọn một mục; lặp 5 lượt. Thử im lặng, lệnh ngoài phạm vi rồi lệnh đúng.',
         'Ghi độ trễ lần đầu/lần sau. Không treo chờ, tự nhận tiếng HOMI hoặc nhảy sai mục; sau lời nhắc vẫn nhận được lệnh đúng.'),
    ]),
    ('3 Dịch tiếng Anh liên tục',
     'Sau khi đổi tài khoản Cloudflare, chạy C01 và D04 trước. Nếu dịch vụ vẫn lỗi/quá hạn mức, ghi Bị chặn và lưu giờ xảy ra.', [
        ('C01', 'Dịch câu mới sau đổi dịch vụ',
         'Vào Dịch; nói “Mình muốn ăn cháo”, “Con để quyển sách màu xanh trên bàn” và thêm một câu tự chọn chưa test.',
         'Câu Việt nhận đúng; bản tiếng Anh đúng nghĩa từng câu. Không trả cùng “Can you say that again, please?” cho mọi câu rõ ràng.'),
        ('C02', 'Mười lượt liên tục',
         'Nói 10 câu khác nhau, mỗi lượt đợi HOMI phát xong. Xen kẽ câu ngắn và dài; nghe đầu, giữa và cuối phiên.',
         'Mỗi lượt có một kết quả; tự mở mic lại đúng lúc. Tiếng/mic luôn qua H20; không lặp câu cũ hoặc mất lượt.'),
        ('C03', 'Ngắt đúng cuối câu',
         'Nói câu khoảng 1–2 giây, rồi câu dài hơn; thử nghỉ nhẹ giữa câu và dừng hẳn cuối câu.',
         'Không cắt mất đầu/đuôi hay ngắt ngay tại nhịp nghỉ ngắn. Câu ngắn không bị kéo thành lượt chờ khoảng 7 giây; ghi độ trễ thực tế.'),
        ('C04', 'Im lặng và tiếng không rõ',
         'Im lặng hai lượt; thử tiếng ồn ngắn, rồi khởi động lại và nói một câu rõ.',
         'Lần đầu nhắc nói lại; lần hai tạm dừng, không mở mic vô hạn. Không tạo bản dịch từ tiếng ồn; lượt hợp lệ sau đó dùng được.'),
        ('C05', 'Dừng rồi MAIN để tiếp tục',
         'Đang dịch, nói đúng “Dừng lại”; chờ. Sau đó nhấn MAIN, chọn “Tiếp tục dịch”.',
         'Chỉ nói “Đã dừng.” và kết thúc phiên. Không tự hỏi tiếp khi chưa nhấn MAIN; sau MAIN mới mở ba lựa chọn và tiếp tục phiên mới.'),
        ('C06', 'Học phần khác khi đang dịch',
         'Đang nghe dịch, nói muốn học phần khác; ở bước chọn, thử tiếp tục dịch rồi thử chuyển sang Chủ đề/Từ vựng.',
         'Không dịch câu điều hướng. Chờ lựa chọn hợp lệ; sang đúng mục, không còn vòng mic hoặc kết quả dịch cũ.'),
        ('C07', 'Mất mạng hoặc lỗi dịch vụ',
         'Ngắt mạng một lượt rồi nối lại. IT mô phỏng lỗi 429/5xx ở môi trường test; không cố làm cạn hạn mức thật.',
         'Không treo, không trả câu tiếng Anh giả như kết quả dịch. Báo lỗi dịch vụ; thử lại được khi đã phục hồi, không tự chuyển Android sang Batch Chunks.'),
    ]),
    ('4 Chủ đề và học từng câu',
     'Chạy một bài đại diện ở mỗi nhóm tuổi và mỗi level. Luồng Core chuẩn là mẫu tiếng Anh → nghĩa tiếng Việt → ting → ghi âm → chấm điểm.', [
        ('D01', 'Vào level và vào bài',
         'Từ Chủ đề chọn level, chủ đề rồi bài chưa học; nghe các phần mở đầu và câu đầu.',
         'Có intro level/bài đúng nội dung, không im lặng hoặc đọc thừa. Tên tiếng Anh nằm trên tiếng Việt; câu đầu đi đủ luồng Core.'),
        ('D02', 'Ting và độ sạch bản ghi',
         'Sau khi nghe mẫu và ting, đọc câu rồi phát lại bản ghi. Lặp 5 lượt, gồm lượt ngay sau cài/mở app.',
         'Mỗi lần mở ghi âm đều có ting; bản ghi không chứa ting/tiếng mẫu, không mất âm đầu. Phát lại đủ nghe, không rè hoặc lúc nhỏ lúc lớn.'),
        ('D03', 'Thời lượng bản ghi',
         'Đọc câu ngắn, câu dài, nghỉ nhẹ giữa câu; thử không nói gì ở một lượt. So độ dài tệp với thời gian thực nói.',
         'Tự dừng theo cuối lời nói, không kéo câu ngắn thành khoảng 7 giây. Không nói thì đợi tối đa khoảng 6 giây; không chấm im lặng là sai.'),
        ('D04', 'Chấm đúng sai bằng API',
         'Với cùng một câu, đọc đúng rõ ràng; lượt khác nói câu sai hẳn; thử lại nói đúng. Ghi giờ của từng lượt.',
         'API phân biệt đúng/sai, phản hồi và tiến độ tương ứng. Không ra cùng kết quả cho mọi lượt; lỗi API không được coi là trẻ nói sai/không rõ.'),
        ('D05', 'MAIN và học câu sau',
         'Trong Core nhấn MAIN, chọn “Câu sau”; thử tương tự bằng nút câu tiếp theo trên màn hình.',
         'Hỏi “Bạn muốn nghe lại, câu trước hay câu sau?”. Nói “Mình học câu sau nhé”, chuyển đúng câu và phát Anh + Việt trước khi ghi.'),
        ('D06', 'Nghe lại và câu trước',
         'Từ câu giữa bài, lần lượt dùng MAIN chọn Nghe lại/Câu trước và nút điều khiển tương ứng.',
         'Chọn đúng câu; phát lại đầy đủ mẫu tiếng Anh và nghĩa Việt rồi mới ghi âm. Không nhảy thẳng vào ghi hoặc tự tiến câu.'),
        ('D07', 'Ranh giới đầu và cuối bài',
         'Ở câu đầu chọn Câu trước; ở câu cuối chọn Câu sau. Kiểm tra cả MAIN và nút chạm khi nút khả dụng.',
         'Đầu: “Đây là câu đầu tiên. Mình nghe lại nhé.” Cuối: “Đây là câu cuối. Bạn hãy hoàn thành câu này nhé.” Sau đó Anh + Việt → mic; không đọc lại tên bài.'),
        ('D08', 'Chấm lần đầu và lưu bản ghi',
         'Chấm 3 lượt ngay sau mở app rồi vài lượt tiếp theo. Ghi nhiều lần cùng câu; mở lịch sử bản ghi và nghe lại.',
         'Không giật, mất lời nhận xét hoặc treo ở chấm điểm. Bản ghi gắn đúng câu/lượt, phát rõ; ghi mới không phá bản ghi và tiến độ đã lưu.'),
    ]),
    ('5 Challenge và học lại',
     'Chạy sau Core. Kiểm tra nội dung thực có trong bài; bài không có nhánh tương ứng thì ghi Không áp dụng.', [
        ('E01', 'Vào Challenge và trả lời',
         'Hoàn thành Core rồi vào Challenge; nghe câu hỏi, chờ ghi âm và trả lời.',
         'Câu nhắc là “Bạn trả lời nhé”. Không ghép thừa đoạn cuối, phát chồng, giật hoặc mất câu hỏi; mic chỉ mở sau lời dẫn.'),
        ('E02', 'Challenge trả lời đúng và sai',
         'Trả lời đúng một câu, sai một câu, im lặng một câu; kiểm tra kết quả cuối phần.',
         'Mỗi lượt chấm/chốt một lần; đúng, sai, không phản hồi được phân biệt. Không cộng sao ngoài quy tắc của Challenge hoặc tính lỗi máy chủ là sai.'),
        ('E03', 'MAIN trong Challenge',
         'Khi đang nghe hoặc chuẩn bị trả lời, nhấn MAIN; chọn Nghe lại, sau đó thử Dừng lại.',
         'Hỏi “Bạn muốn nghe lại hay dừng lại?”. Nghe lại đúng câu hỏi; Dừng lại kết thúc/tạm dừng an toàn, không tiếp tục ghi ngầm.'),
        ('E04', 'Tiếp tục Challenge đã tạm dừng',
         'Tạm dừng Challenge rồi tiếp tục bằng MAIN theo lựa chọn được đọc.',
         'Nói “Mình tiếp tục câu thử thách nhé.” và trở lại đúng câu. Không thêm lời chào/intro bài hay lặp lời thông báo tạm dừng.'),
        ('E05', 'Học tiếp bài đang dở',
         'Dừng giữa bài; đóng/mở HOMI, chọn học tiếp. Thử một bài đã có sao và một bài còn câu chưa đạt.',
         'Giữ đúng tiến độ và câu cần học; không mất hoặc nhân đôi sao. Nếu học lại Core, phát mẫu Anh + Việt trước ghi âm.'),
        ('E06', 'Học lại câu còn thiếu',
         'Mở học lại/ngôi sao còn thiếu trong Chủ đề; chọn một câu, nghe rồi đọc lại. Thử MAIN và Back.',
         'Đúng câu và nghĩa; phát đủ mẫu trước mic, có lời điều hướng. Lượt chấm cập nhật theo quy tắc; Back không kẹt dù HOMI đang nói.'),
        ('E07', 'Bài có bài hát',
         'Vào phần hát sau Challenge; nhấn MAIN và thử Nghe lại, Dừng lại, Bỏ qua. Đối chiếu 5 bài bên dưới.',
         'Bài có Song mới có icon. MAIN hỏi “Bạn muốn nghe lại, dừng lại hay bỏ qua?”. Hết/bỏ bài hát chỉ chuyển tiếp một lần.'),
        ('E08', 'Kết thúc bài và mở khóa',
         'Đi đến cuối bài, cuối chủ đề và cuối level; thử Học tiếp/Học lại/Dừng theo lựa chọn HOMI đọc.',
         'Câu dẫn đúng mốc, lưu tiến độ một lần; mở khóa đúng thứ tự. Không tự nhảy bài/level chưa mở hoặc mất kết quả đã hoàn thành.'),
    ]),
    ('6 Bộ từ vựng',
     'Chuẩn bị một nhóm rỗng và nhóm có dữ liệu; dùng nội dung test riêng cho các thao tác thêm, sửa, xóa.', [
        ('F01', 'Ba nhóm và tìm kiếm',
         'Mở Ba mẹ đã thêm, Ngôi sao của con và Luyện lại; tìm một từ có/không có. So đường vào bằng MAIN và chạm.',
         'Cùng dữ liệu, số lượng đúng; trạng thái rỗng rõ ràng. Ba tiêu đề dùng cùng tone, không có nút Back lồng thừa; tìm kiếm không làm mất mục.'),
        ('F02', 'Thêm nội dung vào hàng chờ',
         'Phụ huynh thêm từ/câu hợp lệ, chọn bản dịch; thử trống, trùng và chọn hơn 3 mục trong một lần thêm.',
         'Mỗi lần tối đa 3 lựa chọn, lưu đúng Anh/Việt. Chặn mục trống/trùng ngoài quy tắc; số mục chờ cập nhật ngay.'),
        ('F03', 'Đầy 5 mục rồi xóa và thêm lại',
         'Đưa hàng chờ đến 5 mục; thử thêm mục thứ 6. Xóa 1–2 mục chưa học rồi thêm lại trong cùng ngày.',
         'Chặn khi đủ 5 mục chờ. Xóa xuống dưới 5 thì thêm được ngay; không dùng tổng số lần từng bấm Thêm để khóa oan.'),
        ('F04', 'Sửa xóa trước và sau khi học',
         'Sửa từ/nghĩa và xóa mục chưa học; bắt đầu học một mục khác rồi thử sửa/xóa lại.',
         'Mục còn chờ sửa/xóa đúng; mục đã bắt đầu học được bảo vệ theo trạng thái. Không mất bản ghi hoặc dữ liệu đã học.'),
        ('F05', 'Danh sách hôm nay của Ba mẹ',
         'Bắt đầu nhóm nội dung mới; nghe vài mục, Back rồi vào lại. Nghe hết nhóm để vào thư viện Ba mẹ.',
         'Từng mục phát Anh → Việt, không mở mic/chấm điểm. Giữ mục chưa nghe; chỉ mở thư viện sau khi nghe đủ nhóm yêu cầu.'),
        ('F06', 'Nghe danh sách và từng mục',
         'Trong Ba mẹ đã thêm/Ngôi sao, bấm loa từng mục và Bắt đầu nghe; thử lựa chọn Mới nhất/Tất cả khi được cung cấp.',
         'Thứ tự đúng lựa chọn; mỗi mục chỉ phát một lần mỗi lượt. Mẫu, nghĩa và bản ghi trẻ đủ nghe, không rè; không tăng âm lượng toàn app ngoài ý muốn.'),
        ('F07', 'MAIN khi học từ vựng',
         'Trong từng nhóm đang học, nhấn MAIN; lần lượt chọn Nghe lại, Câu trước, Câu sau, gồm cả đầu/cuối danh sách.',
         'Hỏi “Bạn muốn nghe lại, câu trước hay câu sau?”. Câu sau có lời dẫn mới; đúng mục, mẫu Anh + Việt trước mic nếu vào lượt luyện nói.'),
        ('F08', 'Luyện lại và nhóm kế tiếp',
         'Chuẩn bị hơn 5 mục luyện lại; học một nhóm gồm đúng/sai và một mục bỏ qua bằng điều hướng. Nhấn MAIN, rồi luyện tiếp.',
         'MAIN có lời dẫn; chốt nhóm một lần. Mục đúng/sai và chưa làm giữ đúng trạng thái; nhóm sau không mất mục hoặc đếm trùng.'),
        ('F09', 'Ngôi sao và dữ liệu sau mở lại',
         'Tạo sao từ một câu Core đủ điều kiện, mở Ngôi sao và nghe bản ghi; đóng/mở HOMI, kiểm tra lại các nhóm.',
         'Đầu nhóm có “Giọng của bạn đây.”, phát đúng bản ghi trẻ. Mục đúng nguồn/nghĩa, không nhân đôi sao; mở lại không mất dữ liệu.'),
    ]),
    ('7 Khóa màn hình và phục hồi phiên học',
     'Không mở hay phát nội dung từ ứng dụng khác. Kiểm tra trên Android thật, ưu tiên máy Xiaomi đang dùng với H20.', [
        ('G01', 'Dịch khi khóa màn hình',
         'Bắt đầu Dịch, khóa màn hình, nói 5 câu; dùng MAIN dừng/tiếp tục rồi mở khóa.',
         'Vẫn nghe và nói qua H20, nhận đúng câu và phát bản dịch. Mở khóa hiển thị kết quả đúng; không rơi sang loa/mic điện thoại.'),
        ('G02', 'Chọn Chủ đề khi khóa màn hình',
         'Khóa màn hình trong phiên hợp lệ; nhấn MAIN, chọn Chủ đề rồi chọn số khi HOMI đọc danh sách.',
         'Cả câu hỏi chọn chủ đề và intro đều qua H20. Vào đúng bài, nghe mẫu rồi ghi/chấm được; không chỉ đổi màn hình mà im tiếng.'),
        ('G03', 'Từ vựng khi khóa màn hình',
         'Từ Dịch và từ menu trang chủ, dùng MAIN vào Từ vựng; thử Ba mẹ đã thêm, Ngôi sao và Luyện lại khi đang khóa.',
         'Vào được từng nhóm, có lời dẫn và MAIN hoạt động. Không mất phản hồi, kẹt lượt hoặc cần mở màn hình mới nghe được.'),
        ('G04', 'Khóa và mở giữa lượt',
         'Khóa/mở màn hình khi đang phát mẫu, ghi và chấm trong Chủ đề/Challenge; lặp 3 vòng.',
         'Không tạo lại câu/lượt chấm, không mất tiếng; trở lại đúng trạng thái. Thông báo phiên nền phù hợp; dừng phiên thì không còn giữ mic.'),
        ('G05', 'Ngắt và kết nối lại H20',
         'Tắt H20 giữa lời dẫn rồi thử giữa ghi âm; bật lại, kết nối và nhấn MAIN. Thử thêm tắt/bật Bluetooth.',
         'Báo trạng thái thật, dừng an toàn; không âm thầm đổi sang điện thoại khi vẫn chọn H20. Kết nối lại đủ nút + mic mới chạy lượt mới.'),
        ('G06', 'Tải offline và quay lại online',
         'Bật/tắt tải dữ liệu offline trong phụ huynh; với gói đã sẵn sàng, ngắt mạng, mở bài đã tải; nối mạng lại và thử Dịch/chấm.',
         'Ghi rõ nội dung nào khả dụng offline. Thiếu gói thì báo rõ, không giả kết quả; trở lại online dùng API được, không treo trạng thái cũ.'),
        ('G07', 'Giao diện và dữ liệu riêng tư',
         'Thử chữ lớn/giao diện tối, Back, lịch sử, đổi tuổi qua cổng phụ huynh. Chỉ trên hồ sơ test, thử hủy rồi xác nhận rút chấp thuận.',
         'Không che MAIN/nút học; dữ liệu hiển thị đúng. Hủy không xóa gì; xác nhận chỉ báo xong khi xử lý thành công, sau đó không tự ghi âm.'),
        ('G08', 'Phiên học liên tục 20 phút',
         'Luân phiên Dịch → Chủ đề → Challenge → Từ vựng; khóa/mở 3 lần. Kết thúc bằng Dừng và mở lại HOMI.',
         'Không crash, đơ, rò mic, lặp lời hoặc tự đổi loa. Dừng xong im tiếng; mở lại còn dữ liệu và có thể bắt đầu phiên mới.'),
    ]),
]


SECTION_ZH = {
    '1 Thiết lập và kết nối H20': ('1 初始设置与 H20 连接', '使用测试资料。不要为了全新安装测试而删除真实应用数据。'),
    '2 MAIN và chuyển giữa các nội dung': ('2 MAIN 与内容切换', '同时测试屏幕 MAIN 与 H20 MAIN。等待 HOMI 播报结束后再回答。'),
    '3 Dịch tiếng Anh liên tục': ('3 连续英语翻译', '更换 Cloudflare 账号后先执行 C01 和 D04。若服务仍受限，请记录“受阻”及发生时间。'),
    '4 Chủ đề và học từng câu': ('4 主题课程与逐句学习', '每个年龄组及每个等级至少执行一课。标准 Core 流程为英语示范、越南语释义、提示音、录音、评分。'),
    '5 Challenge và học lại': ('5 Challenge 与复习', '在 Core 后执行。仅检查课程实际包含的内容；没有相应分支时记录“不适用”。'),
    '6 Bộ từ vựng': ('6 词汇库', '准备一个空分组和一个有数据的分组；新增、编辑和删除只使用测试内容。'),
    '7 Khóa màn hình và phục hồi phiên học': ('7 锁屏与学习会话恢复', '不打开或播放其他应用内容。在 Android 真机上测试，优先使用当前 Xiaomi 手机与 H20。'),
}


# Chinese text is kept next to the Vietnamese source so each row is usable by
# both test teams without switching pages or cross-referencing a second file.
ZH = {
 'A01': ('全新安装与家长确认', '在新测试设备或测试资料中打开 HOMI；阅读法律文件，确认家长身份并同意语音处理。', '按三个设置步骤完成。未阅读或确认前不能开启语音步骤，也不能自动打开麦克风。'),
 'A02': ('拒绝后重新授权', '拒绝麦克风和蓝牙权限并尝试继续；随后重新授权并连接 H20。', '明确提示缺少的权限，不冻结、不显示假连接；重新授权后无需重装即可继续。'),
 'A03': ('选择年龄组', '选择年龄并完成设置后打开主题；分别测试 3–5、6–7、8–10、11–12、13–15 岁。', '目录与课程符合年龄组，重开后仍保存；儿童不能在家长入口外自行修改年龄。'),
 'A04': ('同时连接按键与麦克风', '打开并配对 H20，在 HOMI 中连接；按 MAIN，听完提示后对 H20 麦克风回答。', '只有 BLE 与 H20 音频通道都就绪才显示完成；实际扬声器和麦克风均为 H20。'),
 'A05': ('法律链接与受限模式', '打开隐私政策、使用条款和支持；测试“不使用语音继续”。', '三个链接均可打开；无语音模式不会请求或开启麦克风，家长之后可重新启用。'),
 'A06': ('覆盖安装保留数据', '测试资料已有未完成课程、星星、家长词汇和历史时，覆盖安装待测 APK 并重开。', '年龄、数据和进度均保留且不重复；在文档首页记录实际 APK 和版本。'),
 'B01': ('主页 MAIN', '主页按 MAIN，依次选择英语翻译、主题、词汇库。', '播报三种方向并进入正确年龄内容；每次按键只产生一次提示和一次监听。'),
 'B02': ('双向切换内容', '从翻译、主题、词汇库分别用 MAIN 或语音切到另外两项，共测试六种方向。', '进入新内容前停止旧回合；导航语句不被翻译，不残留旧音频或同时开启两个麦克风。'),
 'B03': ('切换到翻译时有完整开场', '在主题或词汇中说“英语翻译”，然后说一句新内容。', '先播放切换确认，再播“请一句一句说。想停止就说停止”，之后才开麦；不遗漏或重复。'),
 'B04': ('语音选择主题', '经 MAIN 进入主题，用刚播报的编号或名称选择等级、主题和课程；再说无效编号。', '选项与年龄及刚播报列表一致；无效命令只有限次重问，不自动选课。'),
 'B05': ('选择词汇分组', '经 MAIN 进入词汇库，依次选择家长添加、我的星星、复习。', '进入正确分组且与触控数据一致；空分组给出合适提示，不开启无意义录音。'),
 'B06': ('重复 MAIN 与返回键', '快速按 MAIN 三次；在 HOMI 播报、监听、静默时分别按手机返回键。', '不重复菜单或叠音；返回到正确页面并停止旧回合，不冻结或离开后再次播放旧结果。'),
 'B07': ('首次及后续指令识别', '打开应用后用 MAIN 选择内容并重复五次；测试静默、无效命令、再说有效命令。', '记录首次与后续延迟；不长期等待、不把 HOMI 自声当命令、不跳错内容；重试后可识别。'),
 'C01': ('更换服务后的新句翻译', '进入翻译，说“我想喝粥”“孩子把蓝色书放在桌上”及一句未测试的新句。', '越南语识别正确且各句英语含义正确；清晰语音不能全部返回同一句 “Can you say that again, please?”。'),
 'C02': ('连续十轮', '说十个不同句子，每轮等 HOMI 播完；交替短句、长句并检查会话首中尾。', '每轮只有一个结果，播完后正确重开麦；始终使用 H20，不重复旧句或丢轮次。'),
 'C03': ('在句尾正确停止', '说约 1–2 秒短句及较长句；句中短暂停顿，句尾完全停止。', '不切掉开头或结尾，不在句中停顿处结束；短句不会多等到约七秒，并记录实际延迟。'),
 'C04': ('静默与不清楚声音', '连续静默两轮，再测试短噪声；重新开始后清楚说一句。', '首次提示重说，第二次暂停且不无限开麦；噪声不生成翻译，后续有效语句仍可处理。'),
 'C05': ('停止后用 MAIN 继续', '翻译中准确说“Dừng lại”，等待；再按 MAIN 选择继续翻译。', '只播“已停止”并结束会话；未按 MAIN 不继续提问；按后才显示三个选项并开始新会话。'),
 'C06': ('翻译中切换学习内容', '翻译监听时要求学习其他内容；选择继续翻译，再分别切到主题或词汇。', '控制指令不被翻译；等待有效选择后进入正确内容，不残留旧麦克风或翻译结果。'),
 'C07': ('断网或服务错误', '断网一轮后恢复；IT 仅在测试环境模拟 429/5xx，不故意耗尽真实额度。', '不冻结、不把假英语句当翻译结果；显示服务错误，恢复后可重试，Android 不自动回退 Batch Chunks。'),
 'D01': ('进入等级和课程', '从主题选择等级、主题和未学课程；听等级、课程开场和第一句。', '等级与课程开场正确且有声音，不静默、不多播；英文名在越文名上方，第一句走完整 Core。'),
 'D02': ('提示音与录音纯净度', '听完示范与提示音后朗读并回放录音；重复五次，包括刚安装或刚启动后的第一轮。', '每次开麦都有提示音；录音不含提示音或示范，不丢首音；音量足够且不忽大忽小。'),
 'D03': ('录音时长', '朗读短句、长句、句中短暂停顿；一轮保持静默，对比文件时长与实际说话时间。', '按说话结束自动停止，短句不拖成约七秒；静默最多约六秒且不按答错评分。'),
 'D04': ('API 正误评分', '同一句先清楚读对，再说完全错误的句子，最后重试读对；记录每轮时间。', 'API 能区分正误并给出相应反馈和进度；不能所有轮次同结果，API 错误不能算儿童说错或不清楚。'),
 'D05': ('MAIN 学下一句', 'Core 中按 MAIN 选择下一句；再用屏幕下一句按钮重复。', '询问“重听、上一句还是下一句”；播“Mình học câu sau nhé”，切到正确句并先播英越示范再录音。'),
 'D06': ('重听与上一句', '在课程中间分别用 MAIN 和屏幕按钮选择重听及上一句。', '选择正确句子，完整播放英语示范和越南语释义后才录音；不直接开麦或自动跳句。'),
 'D07': ('首句和末句边界', '首句选上一句，末句选下一句；同时测试 MAIN 与可用触控按钮。', '首句提示“这是第一句，再听一次”；末句提示“这是最后一句，请完成”；随后英越示范再开麦，不重播课名。'),
 'D08': ('首次评分与录音保存', '应用启动后立即评分三轮，再继续几轮；同一句多次录音并打开历史回放。', '首次评分不抖动、不丢反馈、不卡住；录音对应正确句子与轮次，新录音不破坏既有记录和进度。'),
 'E01': ('进入 Challenge 并回答', '完成 Core 后进入 Challenge，听问题，等待录音并回答。', '提示为“Bạn trả lời nhé”；无多余尾音、叠音或卡顿，只有提示结束后才开麦。'),
 'E02': ('Challenge 正确、错误与静默', '各回答一题正确、一题错误、一题静默，并查看该部分结果。', '每轮只结算一次并区分三种结果；不额外发 Core 星星，也不把服务器错误算成答错。'),
 'E03': ('Challenge 中 MAIN', '监听或准备回答时按 MAIN；选择重听，再选择停止。', '询问“重听还是停止”；重听正确问题后再开麦；停止后不继续后台录音。'),
 'E04': ('继续已暂停的 Challenge', '暂停 Challenge 后通过 MAIN 选择继续。', '只播“Mình tiếp tục câu thử thách nhé”，再播问题、回答提示并开麦；不增加两个开场。'),
 'E05': ('继续未完成课程', '中途停止并重开 HOMI，选择继续学习；分别测试已有星星和仍未通过的课程。', '保留正确进度和待学句，不丢失或重复星星；重学 Core 时先播英越示范再录音。'),
 'E06': ('复习未获得星星的句子', '在主题中打开复习或缺失星星，选择一句听读；测试 MAIN 与返回键。', '句子与释义正确，录音前完整示范且有导航提示；评分按规则更新，HOMI 播报时返回也不冻结。'),
 'E07': ('含歌曲课程', 'Challenge 后进入歌曲，按 MAIN 并测试重听、停止、跳过；核对下方五首歌。', '只有含歌曲课程显示图标；MAIN 询问“重听、停止还是跳过”；结束或跳过仅前进一次。'),
 'E08': ('课程结束与解锁', '到达课程、主题及等级末尾；按 HOMI 选项测试继续学习、重学、停止。', '各节点提示正确且进度只保存一次；按顺序解锁，不跳到未开放内容，也不丢已完成结果。'),
 'F01': ('三个分组与搜索', '打开家长添加、我的星星、复习；搜索存在和不存在的词，对比 MAIN 与触控入口。', '数据和数量一致，空状态清楚；三处标题色调一致，无多余内层返回键，搜索不丢数据。'),
 'F02': ('新增到等待列表', '家长新增有效词句并选择翻译；测试空值、重复及单次选择超过三项。', '单次最多三项且英越保存正确；无效或重复项按规则拦截，等待数量立即更新。'),
 'F03': ('满五项后删除再新增', '等待列表达到五项后尝试第六项；同日删除 1–2 个未学项再新增。', '满五项时拦截；删到少于五项后立即可新增，不按累计点击新增次数错误锁定。'),
 'F04': ('学习前后编辑删除', '编辑未学项的词和释义并删除；开始学习另一项后再次编辑或删除。', '等待项可正确编辑删除，已开始学习项按状态保护；既有录音和学习数据不丢失。'),
 'F05': ('家长今日列表', '开始一组新内容，听几项后返回再进入；听完整组后进入家长词库。', '每项播放英语再越南语，不开麦、不评分；保留未听项，完成要求后才开放词库。'),
 'F06': ('列表与单项播放', '在家长添加和星星中点单项扬声器及开始播放；若有最新/全部选项则分别测试。', '顺序符合选择，每项每轮只播一次；示范、释义和儿童录音清楚，不意外提高全应用音量。'),
 'F07': ('词汇学习中的 MAIN', '各分组学习中按 MAIN，选择重听、上一句、下一句，并测试列表首尾。', '询问“重听、上一句还是下一句”；下一句使用新提示；若进入口语复习，录音前先播英越示范。'),
 'F08': ('复习及下一组', '准备超过五个复习项；完成一组含正确、错误及一个导航跳过项，按 MAIN 后继续。', 'MAIN 有正确提示且小组只结算一次；正确、错误、未做状态准确，下一组不丢失或重复计数。'),
 'F09': ('星星与重开后的数据', '从符合条件的 Core 句获得星星，打开星星回放录音；重开 HOMI 后检查各组。', '组开始播“Giọng của bạn đây”，播放正确儿童录音；来源和释义正确，不重复星星，重开后数据保留。'),
 'G01': ('锁屏连续翻译', '开始翻译后锁屏，说五句；用 MAIN 停止或继续，再解锁。', '仍通过 H20 听说并正确翻译；解锁显示正确结果，不切到手机扬声器或麦克风。'),
 'G02': ('锁屏选择主题', '在有效会话中锁屏，按 MAIN 选择主题，再按 HOMI 播报选择编号。', '选题问题和开场均从 H20 播放；进入正确课程，示范、录音和评分可用，不会只换页面却静音。'),
 'G03': ('锁屏词汇', '分别从翻译和主页 MAIN 在锁屏时进入词汇，测试家长添加、星星、复习。', '每组均能进入且有提示，MAIN 正常；不丢响应、不卡住，也无需亮屏才听得到。'),
 'G04': ('回合中锁屏与解锁', '主题和 Challenge 在示范、录音、评分阶段分别锁屏解锁，共三轮。', '不重复句子或评分、不丢声音；回到正确状态。后台通知合理，停止会话后不继续占用麦克风。'),
 'G05': ('断开并重连 H20', '提示中关闭 H20，再在录音中重复；重新打开连接并按 MAIN，另测关闭再开启蓝牙。', '显示真实状态并安全停止；选择 H20 时不静默回退手机；按键与麦克风都恢复后才开始新回合。'),
 'G06': ('离线下载与恢复在线', '在家长设置开关离线数据下载；已有包时断网打开已下载课程，恢复网络后测试翻译和评分。', '明确哪些内容可离线；缺包时清楚提示且不伪造结果；恢复在线后 API 可用，不残留旧状态。'),
 'G07': ('界面与隐私数据', '测试大字体、深色、返回、历史及家长入口改年龄；仅测试资料中取消后再确认撤销同意。', 'MAIN 和学习按钮不被遮挡且数据正确；取消不删除，确认仅在服务端处理成功后完成，之后不自动录音。'),
 'G08': ('二十分钟连续会话', '依次使用翻译、主题、Challenge、词汇，锁屏解锁三次；最后停止并重开 HOMI。', '无崩溃、冻结、麦克风泄漏、重复提示或自动换扬声器；停止后静音，重开保留数据并可开始新会话。'),
}


def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement('w:shd')
    node.set(qn('w:fill'), color)
    pr.append(node)


def paragraph(doc, text, style=None, size=None):
    p = doc.add_paragraph(text, style=style)
    if size:
        for r in p.runs:
            r.font.size = Pt(size)
    return p


def add_table(doc, cases):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.15, 7.45, 7.45, 1.70]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for key, val in [('top', '95'), ('bottom', '95'), ('left', '105'), ('right', '105')]:
        node = OxmlElement('w:' + key)
        node.set(qn('w:w'), val)
        node.set(qn('w:type'), 'dxa')
        margins.append(node)
    pr.append(margins)
    headers = ['Mã\n编号', 'Thao tác theo luồng\n流程操作', 'Kết quả cần kiểm tra\n预期结果', 'KQ\n结果\nMã lỗi\n缺陷号']
    header = table.rows[0]
    rep = OxmlElement('w:tblHeader')
    header._tr.get_or_add_trPr().append(rep)
    for i, txt in enumerate(headers):
        header.cells[i].text = txt
        shade(header.cells[i], '17365D')
        for r in header.cells[i].paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string('FFFFFF')
    for idx, (code, title, steps, expected) in enumerate(cases):
        zh_title, zh_steps, zh_expected = ZH[code]
        row = table.add_row()
        row.cells[0].text = code
        p = row.cells[1].paragraphs[0]
        p.add_run(title + '\n').bold = True
        p.add_run(steps + '\n')
        p.add_run(zh_title + '\n').bold = True
        p.add_run(zh_steps)
        p = row.cells[2].paragraphs[0]
        p.add_run(expected + '\n')
        p.add_run(zh_expected)
        row.cells[3].text = '□ Đ 达标\n□ CĐ 未达\n□ B 受阻\n□ NA\n\n______'
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            shade(cell, 'F2F6FA' if idx % 2 else 'FFFFFF')
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.keep_with_next = False
                if i in (0, 3):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9.5)
    return table


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Cm(1.55)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.55)
    sec.right_margin = Cm(1.55)
    sec.footer_distance = Cm(0.65)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.08
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        for el in list(style.element.xpath('.//w:color')):
            for attr in ['themeColor', 'themeTint', 'themeShade']:
                el.attrib.pop(qn('w:' + attr), None)
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Title'].font.bold = True
    title_ppr = doc.styles['Title'].element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn('w:pBdr'))):
        title_ppr.remove(border)
    doc.styles['Heading 1'].font.size = Pt(16)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(4)
    doc.core_properties.title = 'HOMI Kiểm thử Android theo luồng'
    doc.core_properties.subject = 'Test case thiết bị Android và H20'
    doc.core_properties.author = 'HOMI'
    doc.core_properties.keywords = 'Android, H20, kiểm thử, MAIN, Dịch, Chủ đề, Từ vựng'
    count = sum(len(s[2]) for s in SECTIONS)
    paragraph(doc, 'HOMI Kiểm thử Android theo luồng\nHOMI Android 流程测试', 'Title')
    paragraph(doc, f'{count} test case • Ngày 18/09/2026 • Android và H20\n{count} 个测试用例 • 2026年9月18日 • Android 与 H20', size=10)
    paragraph(doc, 'Chạy lần lượt các luồng dưới đây trên điện thoại thật và ghi kết quả sau mỗi ca. Phạm vi không gồm sử dụng HOMI cùng ứng dụng khác.\n请在 Android 真机上按顺序执行以下流程，并在每个用例后记录结果。本范围不包含 HOMI 与其他应用同时使用。')
    paragraph(doc, 'Máy / Android 设备: __________________  Người test 测试人: ______________\nAPK / phiên bản 版本: ________________  Ngày giờ 日期时间: ______________', size=10.5)
    paragraph(doc, 'KQ 结果: Đ = Đạt 达标; CĐ = Chưa đạt 未达标; B = Bị chặn 受阻; NA = Không áp dụng 不适用. Để trống nghĩa là chưa test. Khi lỗi, ghi mã lỗi và giờ xảy ra, kèm ảnh/video hoặc log.\n空白表示尚未测试。发生错误时，请记录缺陷编号、时间，并附图片、视频或日志。', size=10)
    paragraph(doc, 'Chuẩn bị: H20 đủ pin, mạng ổn định, tài khoản dịch/chấm điểm hoạt động. Dùng dữ liệu thử riêng. Lặp Core, Challenge và MAIN trên cả 5 nhóm tuổi; kiểm tra level 1, 2 và 3.\n准备：H20 电量充足、网络稳定、翻译和评分账号可用。只使用测试数据；五个年龄组均执行 Core、Challenge 和 MAIN，并覆盖等级 1、2、3。', size=10)
    for idx, (heading, intro, cases) in enumerate(SECTIONS):
        zh_heading, zh_intro = SECTION_ZH[heading]
        heading_p = paragraph(doc, heading + '\n' + zh_heading, 'Heading 1')
        heading_p.paragraph_format.keep_with_next = True
        intro_p = paragraph(doc, intro + '\n' + zh_intro, size=10.5)
        intro_p.paragraph_format.keep_with_next = True
        add_table(doc, cases)
        if idx == 4:
            paragraph(doc, 'Bài hát cần đối chiếu\n需核对的歌曲', 'Heading 2')
            paragraph(doc, '3–5 tuổi 岁: L1 Numbers / Bài 2 Count With Me; L3 Weather / Bài 2 What to Wear?; L3 My Day / Bài 2 Happy Day.\n6–7 tuổi 岁: L3 My Friends / Bài 1 Play Together.\n8–10 tuổi 岁: L1 My Schedule / Bài 2 My Routine.\nMã bài 歌曲: C35-L1-T02-B02 / C35-L3-T09-B02 / C35-L3-T10-B02 / C67-L3-T08-B01 / C810-L1-T01-B02.', size=10)
        if idx == 6:
            paragraph(doc, 'Tổng kết lượt kiểm thử\n测试汇总', 'Heading 2')
            paragraph(doc, 'Đạt 达标: ____   Chưa đạt 未达: ____   Bị chặn 受阻: ____   NA: ____   Chưa test 未测: ____\nLỗi cần test lại / ghi chú 需复测缺陷 / 备注: __________________________________', size=10.5)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run('HOMI Android  |  ')
    run.font.size = Pt(9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)
    doc.save(TARGET)
    print(f'Created {TARGET}')
    print(f'Test cases: {count}; sections: {len(SECTIONS)}')


if __name__ == '__main__':
    main()
